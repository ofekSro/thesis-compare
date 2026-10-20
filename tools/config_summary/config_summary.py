"""One workbook that gathers everything known about each configuration.

Reads the Phase-1 tables and the production coefficients of one radius-
estimator run and writes outputs\\check_results\\config_summary_<method>.xlsx:

  README          what every column means, the run it came from
  Configs         96 rows: inputs, Pi groups, R_conv,P / R_conv,I measured vs
                  calculated, and a per-config summary of the Z_urban rows
  Z_urban         1920 rows (config x Z_free): R_urban / Z_urban measured vs
                  calculated for pressure and impulse, the five validity
                  conditions of z_urban.z_urban_valid_mask as separate
                  columns, and which of them excluded the row
  Grid_P, Grid_I  config x Z_free matrix of measured Z_urban, grey where the
                  row is outside the model domain, yellow/red by the error of
                  the calculated value

Nothing is recomputed by hand: validity is z_urban_valid_mask, predictions
are the deployed chain (predict_pi / predict_impulse, predict_z_urban,
clip_z_urban_pred). The per-condition columns are a breakdown of that mask
and the tool refuses to write if they do not reproduce it exactly.

Calculated values come from final_production_* (fitted on all 96 configs), so
every error here is IN-SAMPLE. Error statistics use valid rows only; invalid
rows carry a prediction for inspection but never enter an average.

Usage:
    python tools\\config_summary\\config_summary.py
    python tools\\config_summary\\config_summary.py --radius-method req
"""

import argparse
import subprocess
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))  # project root
# Coefficient loaders, error banding and the pressure cross-check are imported
# from the existing tools rather than copied, so the three cannot drift apart.
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'check_formulas'))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'pressure_report'))

import numpy as np
import pandas as pd

from blastlib import constants, paths
from blastlib.processing.radius_estimator import resolve_estimator, VALID_METHODS
from blastlib.regression.convergence_models import predict_impulse, predict_pi
from blastlib.regression.z_urban import (BEYOND_COL, Z_URBAN_ZF_MIN,
                                         clip_z_urban_pred, prepare_maxR_data,
                                         predict_z_urban, z_urban_valid_mask)

from check_formulas import (load_convergence_coefficients,  # noqa: E402
                            load_nonlinear_coefficients)
from pressure_report import (BAND_FILL, DETS, ERR_RED_PCT,  # noqa: E402
                             ERR_YELLOW_PCT, LOCATION, _band_series, _by_det,
                             _check_form, _rel, _rel_err_pct,
                             build_per_z_table)


# (prefix, target_name, Z_urban column, MaxR column, R_conv column)
TARGETS = (('P', 'Pressure', 'Z_urban_P', 'MaxR_P', 'RadiusP'),
           ('I', 'Impulse', 'Z_urban_I', 'MaxR_I', 'RadiusI'))

# The five conditions of z_urban_valid_mask, in its docstring's order.
# Short names are what `excluded_by` lists.
COND_NAMES = ('c1_MaxR<Rconv', 'c2_Rfree<Rconv', 'c3_Rfree>ExcludeR',
              'c4_Zfree>=Zmin', 'c5_Rfree>=b+s')

GREY = ('D9D9D9', '808080')


# ============================================================
# Tables
# ============================================================

def _conditions(df, prefix, target_name, radius_col):
    """The five validity conditions as separate boolean Series (z_urban §mask)."""
    r_free = df['R_free'].astype(float)
    period = df['bsize'].astype(float) + df['swidth'].astype(float)
    return {
        COND_NAMES[0]: ~df[BEYOND_COL[target_name]].astype(bool),
        COND_NAMES[1]: r_free < df[radius_col].astype(float),
        COND_NAMES[2]: r_free > df['exclude_r'].astype(float),
        COND_NAMES[3]: df['Z_free'] >= Z_URBAN_ZF_MIN[target_name],
        COND_NAMES[4]: r_free >= constants.Z_URBAN_MIN_PERIODS * period,
    }


def build_z_table(maxR_df, conv_df, conv_coeffs, z_coeffs):
    """Every (config, Z_free) row with measured vs calculated Z_urban and validity."""
    df = prepare_maxR_data(maxR_df, conv_df).reset_index(drop=True)
    conv_P = _by_det(conv_coeffs, 'RadiusP')
    conv_I = _by_det(conv_coeffs, 'RadiusI')
    W13 = df['weight'].astype(float).values ** (1 / 3)

    out = pd.DataFrame({
        'ConfigName': df['Config'],
        'Det': df['det'].astype(int),
        'Location': df['det'].map(LOCATION),
        'W_kg': df['weight'],
        'b_m': df['bsize'],
        's_m': df['swidth'],
        'H_m': df['height'],
        'rho': df['rho'],
        'Pi_2': df['swidth'] / W13,
        'H_over_s': df['height'] / df['swidth'],
        'Z_free': df['Z_free'],
        'R_free_m': df['R_free'],
        'ExcludeR_m': df['exclude_r'],
        'P_ff_kPa': df['P_ff'],
        'I_ff_kPa_ms': df['I_ff'],
    })

    for prefix, target, zcol, maxr_col, rcol in TARGETS:
        conds = _conditions(df, prefix, target, rcol)
        valid = z_urban_valid_mask(df, zcol, maxr_col, rcol, target)
        both = np.logical_and.reduce([c.values for c in conds.values()])
        if not np.array_equal(both, valid.values):
            raise RuntimeError(
                f'{target}: the five per-condition columns do not reproduce '
                f'z_urban_valid_mask ({int((both != valid.values).sum())} '
                f'rows differ). The mask has changed; update _conditions.')

        z_calc = np.full(len(df), np.nan)
        z_raw = np.full(len(df), np.nan)
        for det in DETS:
            coef = z_coeffs.get((det, target))
            m = (df['det'] == det).values
            if coef is None or not m.any():
                continue
            sub = df[m]
            raw = predict_z_urban(sub['Z_free'].values, sub['rho'].values,
                                  sub['height'].values, sub['swidth'].values,
                                  W13[m], target, coef)
            z_raw[m] = raw
            z_calc[m] = clip_z_urban_pred(raw, sub, det, target,
                                          conv_P, conv_I)

        Zf = df['Z_free'].astype(float).values
        err = _rel_err_pct(df[zcol], z_calc)
        out[f'R_urban_{prefix}_meas_m'] = df[maxr_col]
        out[f'R_urban_{prefix}_calc_m'] = z_calc * W13
        out[f'Z_urban_{prefix}_meas'] = df[zcol]
        out[f'Z_urban_{prefix}_calc'] = z_calc
        out[f'Lambda_{prefix}_meas'] = df[zcol].values / Zf
        out[f'Lambda_{prefix}_calc'] = z_calc / Zf
        out[f'regime_{prefix}'] = np.where(
            df[zcol].values > Zf, 'amplification',
            np.where(df[zcol].values < Zf, 'attenuation',
                     np.where(df[zcol].values == Zf, 'equal', '')))
        out[f'{prefix}_rel_err_pct'] = err
        out[f'{prefix}_err_band'] = np.where(valid.values, _band_series(err), '')
        out[f'{prefix}_clipped_at_Zconv'] = z_calc < z_raw - 1e-12
        out[f'Rconv_{prefix}_meas_m'] = df[rcol]
        for name, c in conds.items():
            out[f'{prefix}_{name}'] = c.values
        out[f'valid_{prefix}'] = valid.values
        names = np.array(list(conds))
        fail = np.column_stack([~c.values for c in conds.values()])
        out[f'excluded_by_{prefix}'] = ['; '.join(names[row]) for row in fail]

    return out.sort_values(['ConfigName', 'Z_free']).reset_index(drop=True)


def build_config_table(conv_df, ztab, conv_coeffs):
    """One row per config: inputs, Pi groups, R_conv measured vs calculated, Z_urban summary."""
    W = conv_df['ChargeWeight'].astype(float).values
    W13 = W ** (1 / 3)
    args = (W, conv_df['AreaDensity'].values, conv_df['Height'].values,
            conv_df['Det'].values, conv_df['StreetWidth'].values,
            conv_df['BuildingSize'].values)
    calc = {'P': predict_pi(*args, _by_det(conv_coeffs, 'RadiusP')),
            'I': predict_impulse(*args, _by_det(conv_coeffs, 'RadiusI'))}

    t = pd.DataFrame({
        'ConfigName': conv_df['ConfigName'],
        'Det': conv_df['Det'].astype(int),
        'Location': conv_df['Det'].map(LOCATION),
        'W_kg': W,
        'b_m': conv_df['BuildingSize'],
        's_m': conv_df['StreetWidth'],
        'H_m': conv_df['Height'],
        'rho': conv_df['AreaDensity'],
        'Pi_2': conv_df['StreetWidth'] / W13,
        'H_over_s': conv_df['Height'] / conv_df['StreetWidth'],
        'VolumeDensity': conv_df['VolumeDensity'],
        'W13': W13,
    })
    for prefix, _, _, _, rcol in TARGETS:
        meas = conv_df[rcol].astype(float).values
        t[f'Rconv_{prefix}_meas_m'] = meas
        t[f'Rconv_{prefix}_calc_m'] = calc[prefix]
        t[f'Rconv_{prefix}_rel_err_pct'] = _rel_err_pct(meas, calc[prefix])
        t[f'Rconv_{prefix}_err_band'] = _band_series(
            t[f'Rconv_{prefix}_rel_err_pct'])
        t[f'Zconv_{prefix}_meas'] = meas / W13
        t[f'Zconv_{prefix}_calc'] = calc[prefix] / W13
    t['PressureAtR_kPa'] = conv_df['PressureAtR']
    t['ImpulseAtR_kPa_ms'] = conv_df['ImpulseAtR']

    idx = t['ConfigName']
    for prefix, *_ in TARGETS:
        v = ztab[ztab[f'valid_{prefix}']]
        g = v.groupby('ConfigName')
        abs_err = v[f'{prefix}_rel_err_pct'].abs().groupby(v['ConfigName'])
        t[f'n_Z_valid_{prefix}'] = idx.map(g.size()).fillna(0).astype(int)
        t[f'Zfree_valid_min_{prefix}'] = idx.map(g['Z_free'].min())
        t[f'Zfree_valid_max_{prefix}'] = idx.map(g['Z_free'].max())
        t[f'Zurban_{prefix}_MAPE_pct'] = idx.map(abs_err.mean())
        t[f'Zurban_{prefix}_MAPE_band'] = _band_series(
            t[f'Zurban_{prefix}_MAPE_pct'])
        t[f'Zurban_{prefix}_worst_abs_err_pct'] = idx.map(abs_err.max())
    return t


def build_grid(ztab, prefix):
    """config x Z_free matrices: measured Z_urban, validity and error band."""
    piv = lambda col: ztab.pivot(index='ConfigName', columns='Z_free',
                                 values=col)
    return (piv(f'Z_urban_{prefix}_meas'), piv(f'valid_{prefix}'),
            piv(f'{prefix}_err_band'))


def cross_check_pressure(ztab, maxR_df, conv_df, conv_coeffs, z_coeffs):
    """Max |difference| against tools/pressure_report on its valid pressure rows."""
    ref = build_per_z_table(maxR_df, conv_df, conv_coeffs, z_coeffs)
    mine = ztab[ztab['valid_P']].set_index(['ConfigName', 'Z_free'])
    ref = ref.set_index(['ConfigName', 'Z_free'])
    if set(mine.index) != set(ref.index):
        raise RuntimeError('Valid pressure rows differ from pressure_report.')
    d = (mine.loc[ref.index, 'R_urban_P_calc_m']
         - ref['MaxR_P_calculated_m']).abs().max()
    if not d < 1e-9:
        raise RuntimeError(f'Pressure R_urban differs from pressure_report '
                           f'by up to {d:g} m.')
    return len(ref)


# ============================================================
# Excel output
# ============================================================

NUM_FORMATS = (('_pct', '+0.00;-0.00'), ('_m', '0.00'), ('rho', '0.0000'),
               ('Pi_2', '0.000'), ('H_over_s', '0.000'), ('W13', '0.000'),
               ('VolumeDensity', '0.0000'), ('_meas', '0.000'),
               ('_calc', '0.000'), ('_kPa', '0.00'), ('_kPa_ms', '0.00'))

BANDED = {
    'Configs': [(f'Rconv_{p}_err_band', (f'Rconv_{p}_calc_m',
                                         f'Rconv_{p}_rel_err_pct'))
                for p in 'PI'] +
               [(f'Zurban_{p}_MAPE_band', (f'Zurban_{p}_MAPE_pct',))
                for p in 'PI'],
    'Z_urban': [(f'{p}_err_band', (f'Z_urban_{p}_calc',
                                   f'{p}_rel_err_pct')) for p in 'PI'],
}


def _fmt(col):
    for suffix, f in NUM_FORMATS:
        if col == suffix or col.endswith(suffix):
            return f
    return None


def _style_table(ws, df, bands, grey_col=None):
    """Header, freeze, filter, number formats, band colours, grey invalid rows."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    cols = list(df.columns)
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = ws.dimensions
    for i, name in enumerate(cols, start=1):
        h = ws.cell(row=1, column=i)
        h.font = Font(bold=True)
        h.alignment = Alignment(horizontal='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = max(
            10, min(len(name) + 2, 24))
        f = _fmt(name)
        if f:
            for r in range(2, len(df) + 2):
                ws.cell(row=r, column=i).number_format = f
    ws.column_dimensions['A'].width = 34

    fill = lambda c: PatternFill('solid', start_color=c, end_color=c)
    for band_col, painted in bands:
        targets = [cols.index(c) + 1 for c in painted if c in cols]
        for r, band in enumerate(df[band_col].values, start=2):
            if band in BAND_FILL:
                bg, fg = BAND_FILL[band]
                for c in targets:
                    ws.cell(row=r, column=c).fill = fill(bg)
                    ws.cell(row=r, column=c).font = Font(color=fg)
    if grey_col:
        # Invalid rows: grey the calculated/error cells of that target only.
        for prefix in 'PI':
            targets = [i + 1 for i, c in enumerate(cols)
                       if c in (f'R_urban_{prefix}_calc_m',
                                f'Z_urban_{prefix}_calc',
                                f'Lambda_{prefix}_calc',
                                f'{prefix}_rel_err_pct')]
            for r, ok in enumerate(df[f'valid_{prefix}'].values, start=2):
                if not ok:
                    for c in targets:
                        ws.cell(row=r, column=c).fill = fill(GREY[0])
                        ws.cell(row=r, column=c).font = Font(color=GREY[1])


def _write_grid(writer, name, values, valid, band):
    """One config x Z_free sheet, coloured cell by cell."""
    from openpyxl.styles import Font, PatternFill
    values.to_excel(writer, sheet_name=name)
    ws = writer.sheets[name]
    ws.freeze_panes = 'B2'
    ws.column_dimensions['A'].width = 34
    fill = lambda c: PatternFill('solid', start_color=c, end_color=c)
    for i in range(values.shape[0]):
        for j in range(values.shape[1]):
            cell = ws.cell(row=i + 2, column=j + 2)
            cell.number_format = '0.00'
            if not bool(valid.iat[i, j]):
                cell.fill = fill(GREY[0])
                cell.font = Font(color=GREY[1])
            elif band.iat[i, j] in BAND_FILL:
                bg, fg = BAND_FILL[band.iat[i, j]]
                cell.fill = fill(bg)
                cell.font = Font(color=fg)


def _git_rev():
    try:
        rev = subprocess.run(['git', 'rev-parse', '--short', 'HEAD'],
                             cwd=paths.PROJECT_ROOT, capture_output=True,
                             text=True, check=True).stdout.strip()
        # --no-optional-locks: never rewrite .git/index (OneDrive syncs it).
        dirty = subprocess.run(['git', '--no-optional-locks', 'status',
                                '--porcelain', '--',
                                'blastlib', 'tools', 'outputs/tables'],
                               cwd=paths.PROJECT_ROOT, capture_output=True,
                               text=True).stdout.strip()
        return rev + (' (uncommitted changes in code/tables)' if dirty else '')
    except Exception:
        return 'unknown'


def _readme(method, n_cfg, n_rows, n_check):
    rows = [
        ('Radius estimator', method),
        ('Commit', _git_rev()),
        ('Written', date.today().isoformat()),
        ('Sources', f'convergence_table_{method}.csv, max_radius_per_Z_{method}'
                    f'.csv, final_production_*_{method}.csv'),
        ('IN-SAMPLE', 'Calculated values use the production coefficients, '
                      'fitted on all 96 configs. Every error is in-sample. '
                      r'Held-out errors: tools\check_formulas.'),
        ('Rows', f'{n_cfg} configs, {n_rows} (config, Z_free) rows'),
        ('Cross-check', f'{n_check} valid pressure rows match '
                        'tools/pressure_report exactly'),
        ('', ''),
        ('--- Configs ---', ''),
        ('W_kg, b_m, s_m, H_m', 'Charge weight [kg TNT], building footprint '
                                'side, street width, building height [m]'),
        ('rho', 'b^2/(b+s)^2, plan area density [-]'),
        ('Pi_2', 's/W^(1/3) [m/kg^(1/3)]'),
        ('H_over_s', 'H/s, canyon aspect ratio [-]'),
        ('W13', 'W^(1/3) [kg^(1/3)]'),
        ('Rconv_X_meas_m / _calc_m', 'Convergence radius, measured (Phase 1) '
                                     'and calculated (production fit) [m]; '
                                     'X = P pressure, I impulse'),
        ('Zconv_X', 'Rconv / W^(1/3)'),
        ('PressureAtR / ImpulseAtR', 'Free-field level at R_conv,P / R_conv,I'),
        ('n_Z_valid_X', 'Number of Z_free levels inside the Z_urban domain'),
        ('Zurban_X_MAPE_pct', 'Mean |relative error| of calculated Z_urban '
                              'over the valid rows of that config'),
        ('', ''),
        ('--- Z_urban ---', ''),
        ('Z_free, R_free_m', 'Free-field scaled distance and R_free = '
                             'Z_free*W^(1/3) [m]'),
        ('P_ff_kPa, I_ff_kPa_ms', 'Free-field peak pressure / impulse at '
                                  'Z_free'),
        ('R_urban_X_meas_m', 'MaxR: farthest distance from the charge at '
                             'which the urban field still delivers that '
                             'free-field level [m]'),
        ('Z_urban_X', 'R_urban / W^(1/3); _calc is the deployed chain, '
                      'clipped from above at the predicted Z_conv'),
        ('Lambda_X', 'Z_urban / Z_free (>1 amplification, <1 attenuation)'),
        ('X_clipped_at_Zconv', 'True where the Z_conv clip changed the '
                               'prediction'),
        ('c1..c5 columns', 'The five conditions of '
                           'z_urban.z_urban_valid_mask: c1 MaxR inside '
                           'R_conv (beyond flag), c2 R_free < R_conv, c3 '
                           'R_free > ExcludeR (first street), c4 Z_free >= 2,'
                           f' c5 R_free >= {constants.Z_URBAN_MIN_PERIODS:g}'
                           '*(b+s) (D46)'),
        ('valid_X', 'All five conditions hold: the row is in the domain the '
                    'Z_urban model is fitted and used on'),
        ('excluded_by_X', 'Conditions that fail'),
        ('', ''),
        ('--- Colours ---', ''),
        ('Grey', 'Row outside the domain; prediction shown, not in any '
                 'average'),
        ('Yellow', f'|relative error| {ERR_YELLOW_PCT:g}-{ERR_RED_PCT:g}%'),
        ('Red', f'|relative error| > {ERR_RED_PCT:g}%'),
        ('Grid_P / Grid_I', 'Measured Z_urban per config x Z_free, coloured '
                            'by validity and by the error of the calculated '
                            'value'),
    ]
    return pd.DataFrame(rows, columns=['Item', 'Meaning'])


def write_workbook(path, readme, configs, ztab, grids):
    from openpyxl.styles import Alignment, Font
    with pd.ExcelWriter(path, engine='openpyxl') as w:
        readme.to_excel(w, sheet_name='README', index=False)
        configs.to_excel(w, sheet_name='Configs', index=False)
        ztab.to_excel(w, sheet_name='Z_urban', index=False)
        for prefix, (vals, valid, band) in grids.items():
            _write_grid(w, f'Grid_{prefix}', vals, valid, band)
        _style_table(w.sheets['Configs'], configs, BANDED['Configs'])
        _style_table(w.sheets['Z_urban'], ztab, BANDED['Z_urban'],
                     grey_col=True)
        ws = w.sheets['README']
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 100
        for r in range(1, len(readme) + 2):
            ws.cell(row=r, column=1).font = Font(bold=True)
            ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True,
                                                           vertical='top')
    return path


# ============================================================
# Entry points
# ============================================================

def main(*, tables_dir=None, out_dir=None, radius_method=None,
         progress=print):
    """Build the summary workbook. Returns a dict with counts and the output path."""
    method = resolve_estimator(radius_method)['method']
    tables = paths.resolve(tables_dir, paths.TABLES_DIR)
    out = paths.ensure_dir(paths.resolve(out_dir, paths.CHECK_RESULTS_DIR))

    conv_csv = paths.conv_csv(method, tables)
    maxr_csv = paths.maxr_csv(method, tables)
    conv_coef_csv = tables / paths.suffixed(
        'final_production_convergence_coefficients.csv', method)
    z_coef_csv = tables / paths.suffixed(
        'final_production_z_urban_coefficients.csv', method)
    for p in (conv_csv, maxr_csv, conv_coef_csv, z_coef_csv):
        if not Path(p).is_file():
            raise FileNotFoundError(f'Missing {p}. Run run_analysis.py '
                                    f'--phase all --radius-method {method}')

    conv_df = pd.read_csv(conv_csv)
    maxR_df = pd.read_csv(maxr_csv)
    conv_coeffs = load_convergence_coefficients(conv_coef_csv)
    z_coeffs = load_nonlinear_coefficients(z_coef_csv)
    for target in ('Pressure', 'Impulse'):
        _check_form(z_coeffs, target)

    ztab = build_z_table(maxR_df, conv_df, conv_coeffs, z_coeffs)
    configs = build_config_table(conv_df, ztab, conv_coeffs)
    n_check = cross_check_pressure(ztab, maxR_df, conv_df, conv_coeffs,
                                   z_coeffs)
    grids = {p: build_grid(ztab, p) for p in 'PI'}

    xlsx = out / paths.suffixed('config_summary.xlsx', method)
    write_workbook(xlsx, _readme(method, len(configs), len(ztab), n_check),
                   configs, ztab, grids)

    progress(f'CONFIG SUMMARY - radius method {method!r} (in-sample)')
    progress(f'  {len(configs)} configs, {len(ztab)} (config, Z) rows')
    for p in 'PI':
        v = ztab[ztab[f'valid_{p}']]
        progress(f'  {p}: R_conv MAPE '
                 f'{configs[f"Rconv_{p}_rel_err_pct"].abs().mean():.2f}%   '
                 f'valid Z rows {len(v)}   Z_urban MAPE '
                 f'{v[f"{p}_rel_err_pct"].abs().mean():.2f}%')
    progress(f'  pressure cross-check vs pressure_report: {n_check} rows '
             f'identical')
    progress(f'Saved: {_rel(xlsx)}')
    return {'method': method, 'n_configs': len(configs),
            'n_rows': len(ztab), 'xlsx': str(xlsx)}


def cli(argv=None):
    p = argparse.ArgumentParser(
        description='Per-config summary workbook: inputs, Pi groups, '
                    'convergence radii and Z_urban, measured vs calculated.')
    p.add_argument('--tables-dir', default=None)
    p.add_argument('--out-dir', default=None,
                   help='Output folder (default outputs/check_results).')
    p.add_argument('--radius-method', default=None, dest='radius_method',
                   help=f'{"|".join(VALID_METHODS)}, pXX or a soft token '
                        "like 'req_soft3' (default: constants.RADIUS_ESTIMATOR).")
    a = p.parse_args(argv)
    return main(tables_dir=a.tables_dir, out_dir=a.out_dir,
                radius_method=a.radius_method)


if __name__ == '__main__':
    cli()
